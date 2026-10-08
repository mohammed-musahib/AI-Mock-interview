import os
import re
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt'}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

COMMON_TECH_SKILLS = [
    "python", "flask", "django", "fastapi", "java", "spring boot", "spring", "hibernate",
    "javascript", "typescript", "react", "next.js", "vue", "angular", "node.js", "express",
    "html5", "css3", "tailwind", "bootstrap", "sql", "postgresql", "mysql", "mongodb", "redis",
    "docker", "kubernetes", "aws", "azure", "gcp", "git", "ci/cd", "linux", "rest api", "graphql",
    "machine learning", "deep learning", "nlp", "computer vision", "tensorflow", "pytorch",
    "scikit-learn", "pandas", "numpy", "data analysis", "power bi", "tableau", "spark",
    "cyber security", "penetration testing", "ethical hacking", "siem", "soc", "wireshark",
    "metasploit", "cryptography", "network security", "owasp", "incident response", "burp suite"
]

def allowed_file(filename: str) -> bool:
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(file_path: str) -> str:
    text_content = []
    reader = PdfReader(file_path)
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            text_content.append(extracted)
    return "\n".join(text_content)

def extract_text_from_docx(file_path: str) -> str:
    doc = Document(file_path)
    text_content = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_content.append(" | ".join(row_text))
    return "\n".join(text_content)

def extract_text_from_txt(file_path: str) -> str:
    encodings = ['utf-8', 'latin-1', 'cp1252']
    for enc in encodings:
        try:
            with open(file_path, 'r', encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()

def clean_extracted_text(raw_text: str) -> str:
    if not raw_text:
        return ""
    # Normalize line breaks and remove repetitive whitespace
    text = re.sub(r'\r\n|\r', '\n', raw_text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def detect_skills_from_text(text: str) -> list[str]:
    lower_text = text.lower()
    matched = []
    for skill in COMMON_TECH_SKILLS:
        pattern = r'\b' + re.escape(skill) + r'\b'
        if re.search(pattern, lower_text):
            # Format nicely for display
            matched.append(skill.title() if len(skill) > 4 else skill.upper())
    return sorted(list(set(matched)))[:15]

def process_resume_file(uploaded_file, upload_folder: str) -> dict:
    if not uploaded_file or uploaded_file.filename == '':
        return {"success": False, "error": "No file selected."}

    if not allowed_file(uploaded_file.filename):
        return {
            "success": False,
            "error": "Unsupported file format. Please upload PDF, DOCX, or TXT."
        }

    original_name = uploaded_file.filename
    clean_name = secure_filename(original_name)
    if not clean_name:
        clean_name = "resume_file"

    file_ext = clean_name.rsplit('.', 1)[1].lower() if '.' in clean_name else "txt"
    temp_path = os.path.join(upload_folder, f"temp_{os.urandom(8).hex()}.{file_ext}")

    try:
        os.makedirs(upload_folder, exist_ok=True)
        uploaded_file.save(temp_path)

        file_size = os.path.getsize(temp_path)
        if file_size > MAX_FILE_SIZE_BYTES:
            return {
                "success": False,
                "error": "File size exceeds the 10 MB limit. Please upload a smaller resume."
            }

        if file_ext == 'pdf':
            raw_text = extract_text_from_pdf(temp_path)
        elif file_ext == 'docx':
            raw_text = extract_text_from_docx(temp_path)
        else:
            raw_text = extract_text_from_txt(temp_path)

        cleaned_text = clean_extracted_text(raw_text)

        if not cleaned_text or len(cleaned_text) < 50:
            return {
                "success": False,
                "error": "We could not extract readable text from this file. Please verify it is not an image-only scan."
            }

        detected_skills = detect_skills_from_text(cleaned_text)
        word_count = len(cleaned_text.split())

        return {
            "success": True,
            "filename": original_name,
            "text": cleaned_text,
            "char_count": len(cleaned_text),
            "word_count": word_count,
            "skills": detected_skills
        }
    except Exception as e:
        return {
            "success": False,
            "error": f"Failed to process file: {str(e)}"
        }
    finally:
        # Transient storage cleanup for safety and privacy
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass
