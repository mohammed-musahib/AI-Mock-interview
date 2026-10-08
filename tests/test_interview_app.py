import io
import pytest
from app import app
from pypdf import PdfWriter
from docx import Document

@pytest.fixture
def client():
    app.config["TESTING"] = True
    app.config["SECRET_KEY"] = "test_secret_key"
    with app.test_client() as client:
        yield client

def create_sample_pdf_bytes():
    from reportlab.pdfgen import canvas
    packet = io.BytesIO()
    can = canvas.Canvas(packet)
    can.drawString(10, 50, "Full Stack Developer with Python, Flask, React, Docker, and PostgreSQL.")
    can.drawString(10, 30, "Built scalable microservices and REST APIs with Redis caching.")
    can.save()
    packet.seek(0)
    return packet.getvalue()

def create_sample_docx_bytes():
    doc = Document()
    doc.add_heading("Senior Security Engineer Resume", 0)
    doc.add_paragraph("Specialized in Cyber Security, OWASP, Penetration Testing, SIEM, and SOC operations.")
    packet = io.BytesIO()
    doc.save(packet)
    packet.seek(0)
    return packet.getvalue()

def test_landing_page(client):
    response = client.get("/")
    assert response.status_code == 200
    assert b"AI Mock Interview" in response.data

def test_setup_page(client):
    response = client.get("/setup")
    assert response.status_code == 200
    assert b"Select Interview Domain" in response.data

def test_upload_txt_resume(client):
    resume_content = b"John Doe\nSoftware Engineer with experience in Python, Flask, SQL, Docker, and AWS."
    data = {
        "resume": (io.BytesIO(resume_content), "resume.txt")
    }
    response = client.post("/api/upload-resume", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    json_data = response.get_json()
    assert json_data["success"] is True
    assert "Python" in json_data["skills"] or "Flask" in json_data["skills"]

def test_upload_pdf_resume(client):
    pdf_bytes = create_sample_pdf_bytes()
    data = {
        "resume": (io.BytesIO(pdf_bytes), "sample_resume.pdf")
    }
    response = client.post("/api/upload-resume", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    json_data = response.get_json()
    assert json_data["success"] is True
    assert json_data["filename"] == "sample_resume.pdf"

def test_upload_docx_resume(client):
    docx_bytes = create_sample_docx_bytes()
    data = {
        "resume": (io.BytesIO(docx_bytes), "sample_resume.docx")
    }
    response = client.post("/api/upload-resume", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    json_data = response.get_json()
    assert json_data["success"] is True

def test_upload_invalid_extension(client):
    data = {
        "resume": (io.BytesIO(b"fake image data"), "avatar.png")
    }
    response = client.post("/api/upload-resume", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    json_data = response.get_json()
    assert json_data["success"] is False

def test_generate_interview_without_resume(client):
    response = client.post("/api/generate-interview", json={"domain": "Python Full Stack"})
    assert response.status_code == 400
    json_data = response.get_json()
    assert "upload your resume" in json_data["error"]

def test_generate_interview_invalid_domain(client):
    client.post("/api/upload-resume", data={"resume": (io.BytesIO(b"Resume content..."), "dev.txt")}, content_type="multipart/form-data")
    response = client.post("/api/generate-interview", json={"domain": "Invalid Domain XYZ"})
    assert response.status_code == 400

@pytest.mark.parametrize("domain", [
    "Cyber Security",
    "Python Full Stack",
    "Java Full Stack",
    "Data Science",
    "AI/ML Engineer"
])
def test_all_five_domains_generation(client, domain):
    # Upload resume
    resume_text = f"Candidate experienced in {domain} and associated system architectures."
    client.post("/api/upload-resume", data={"resume": (io.BytesIO(resume_text.encode()), "resume.txt")}, content_type="multipart/form-data")

    # Generate
    res = client.post("/api/generate-interview", json={"domain": domain})
    assert res.status_code == 200
    data = res.get_json()
    assert data["success"] is True
    assert len(data["questions"]) == 10
    assert data["domain"] == domain

def test_full_interview_lifecycle(client):
    # 1. Upload Resume
    resume_content = b"Candidate Profile: Java, Spring Boot, MySQL, REST API, Microservices."
    client.post("/api/upload-resume", data={"resume": (io.BytesIO(resume_content), "java_dev.txt")}, content_type="multipart/form-data")

    # 2. Generate Interview for Java Full Stack
    gen_response = client.post("/api/generate-interview", json={"domain": "Java Full Stack"})
    assert gen_response.status_code == 200
    gen_data = gen_response.get_json()
    assert gen_data["success"] is True
    assert gen_data["domain"] == "Java Full Stack"
    questions = gen_data["questions"]
    assert len(questions) == 10

    # Ensure questions 1-5 are MCQ and do NOT leak correct_answer
    for i in range(5):
        assert questions[i]["type"] == "mcq"
        assert len(questions[i]["options"]) == 4
        assert "correct_answer" not in questions[i]

    # Ensure questions 6-10 are theory
    for i in range(5, 10):
        assert questions[i]["type"] == "theory"

    # 3. Interview page is accessible
    interview_page = client.get("/interview")
    assert interview_page.status_code == 200

    # 4. Submit Answers
    answers = {
        "1": "B",
        "2": "B",
        "3": "B",
        "4": "B",
        "5": "B",
        "6": "Java Memory Model separates Heap for object allocations and Stack for thread call frames and local references. OutOfMemoryError is diagnosed with heap dumps.",
        "7": "@Transactional manages boundary commit and rollbacks for checked and unchecked exceptions.",
        "8": "Microservices communicate via Spring Cloud Gateway, Feign clients, and Resilience4j circuit breakers.",
        "9": "Virtual Threads enable lightweight threads scheduled by the JVM on carrier OS threads.",
        "10": "JUnit 5 and Mockito for unit testing with Testcontainers for integration database verification."
    }

    eval_response = client.post("/api/evaluate-interview", json={"answers": answers})
    assert eval_response.status_code == 200
    eval_data = eval_response.get_json()
    assert eval_data["success"] is True
    results = eval_data["results"]

    # Verify score scale is 1 to 10
    final_score = results["final_score"]
    assert 1.0 <= final_score <= 10.0
    assert results["mcq_score"]["total"] == 5
    assert results["theory_score"]["total"] == 5.0
    assert "report" in results
    assert len(results["reviews"]) == 10

    # 5. Results page is accessible
    results_page = client.get("/results")
    assert results_page.status_code == 200
    assert b"Your Interview" in results_page.data

    # 6. Session reset works
    reset_res = client.post("/api/reset-interview")
    assert reset_res.status_code == 200
